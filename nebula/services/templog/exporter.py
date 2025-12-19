"""Excel export functionality for temperature logs"""

from pathlib import Path
from typing import List, Dict, Optional
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

from nebula.observability import get_logger

logger = get_logger(__name__)


class TempLogExporter:
    """Export temperature logs to Excel with color-coding"""

    def __init__(self, output_dir: Optional[Path] = None):
        """
        Initialize exporter

        Args:
            output_dir: Directory for exported files
        """
        if output_dir is None:
            output_dir = Path(__file__).parent.parent.parent / "exports"

        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        entries: List[Dict],
        date: str,
        unit: Optional[str] = None
    ) -> Path:
        """
        Export temperature log entries to Excel file

        Args:
            entries: List of entry dictionaries
            date: Date for filename
            unit: Optional unit for filename

        Returns:
            Path to generated Excel file
        """
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Temperature Log"

        # Define headers
        headers = [
            "Date",
            "Time",
            "Unit",
            "Station",
            "Item",
            "Temp (°F)",
            "Status",
            "Corrective Action",
            "Initials",
            "Notes"
        ]

        # Write headers
        ws.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Define status colors
        status_fills = {
            "PASS": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),  # Green
            "FAIL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),  # Red
            "REVIEW": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Yellow
        }

        # Write data rows
        for entry in entries:
            row_data = [
                entry.get("date", ""),
                entry.get("time", ""),
                entry.get("unit", ""),
                entry.get("station", ""),
                entry.get("item", ""),
                entry.get("temp_f", ""),
                entry.get("status", ""),
                entry.get("corrective_action", ""),
                entry.get("initials", ""),
                entry.get("notes", "")
            ]
            ws.append(row_data)

            # Color-code status cell
            status = entry.get("status")
            if status in status_fills:
                # Status is column 7 (G)
                status_cell = ws.cell(row=ws.max_row, column=7)
                status_cell.fill = status_fills[status]

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Generate filename
        unit_part = f"_{unit}" if unit else ""
        filename = f"{date}{unit_part}_temp_log.xlsx"
        output_path = self.output_dir / filename

        # Save workbook
        wb.save(output_path)
        logger.info(f"Exported {len(entries)} entries to {output_path}")

        return output_path
